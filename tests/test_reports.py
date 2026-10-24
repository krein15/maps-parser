"""Отчёты: книга Excel и выгрузка CSV."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from mapsparser import csvout, excel
from mapsparser.gaps import Context, apply
from mapsparser.models import EMAIL, PHONE, WEBSITE, Company, Link, Place
from mapsparser.pipeline import Result, Settings
from mapsparser.sources import TwoGis, Yandex


def sample_result(tyumen) -> Result:
    rich = Place(
        name="Ола Ола",
        address="Профсоюзная улица, 56",
        city="Тюмень",
        lat=57.152185,
        lon=65.554530,
        rubrics=("Кофейни",),
        by_source={
            "2gis": Company(
                source="2gis", source_id="1", name="Ола Ола", url="https://2gis.ru/tyumen/firm/1",
                address="Профсоюзная улица, 56", rating=4.8, reviews=125, schedule="Пн–Пт 09:00–22:00",
                links=(Link(PHONE, "+73452666707"), Link(EMAIL, "ola@mail.ru"),
                       Link(WEBSITE, "https://ola-ola.ru"), Link("telegram", "https://t.me/olaola")),
            ),
            "yandex": Company(
                source="yandex", source_id="2", name="Ола Ола", url="https://yandex.ru/maps/org/ola_ola/2/",
                address="Профсоюзная улица, 56", rating=5.0, reviews=130, verified_owner=True, photos=81,
            ),
        },
    )
    poor = Place(
        name="Кофе и Точка",
        address="улица Ленина, 57А",
        city="Тюмень",
        by_source={
            "2gis": Company(
                source="2gis", source_id="3", name="Кофе и Точка", url="https://2gis.ru/tyumen/firm/3",
                address="улица Ленина, 57А", rating=3.6, reviews=4,
                links=(Link(PHONE, "+73452000000"),),
            )
        },
    )
    context = Context(sources=(TwoGis, Yandex), exhaustive=frozenset({"2gis", "yandex"}))
    places = apply([rich, poor], context)
    return Result(
        settings=Settings(query="кофейня", city=tyumen, sources=[TwoGis, Yandex]),
        places=places,
        started_at=datetime(2026, 10, 25, 12, 0).astimezone(),
        finished_at=datetime(2026, 10, 25, 12, 4).astimezone(),
        context=context,
    )


@pytest.fixture
def report(tyumen) -> Result:
    return sample_result(tyumen)


def test_excel_has_both_sheets(report, tmp_path: Path) -> None:
    path = excel.write(report, tmp_path / "отчёт.xlsx")
    book = load_workbook(path)
    assert book.sheetnames == ["Фирмы", "Сводка"]


def test_excel_row_holds_the_contacts_and_the_gaps(report, tmp_path: Path) -> None:
    sheet = load_workbook(excel.write(report, tmp_path / "отчёт.xlsx"))["Фирмы"]
    titles = [cell.value for cell in sheet[1]]
    assert titles[:4] == ["Фирма", "Чего нет", "Оценка", "Чем заняться"]
    rows = {row[0]: row for row in sheet.iter_rows(min_row=2, values_only=True)}
    ola = rows["Ола Ола"]
    assert "ola@mail.ru" in ola[titles.index("Почта")]
    assert "+7 (345) 266-67-07" in ola[titles.index("Телефон")]
    assert "Telegram" in ola[titles.index("Соцсети и мессенджеры")]
    assert "нет" in (rows["Кофе и Точка"][1] or "")


def test_excel_links_to_the_card_in_each_map(report, tmp_path: Path) -> None:
    sheet = load_workbook(excel.write(report, tmp_path / "отчёт.xlsx"))["Фирмы"]
    titles = [cell.value for cell in sheet[1]]
    column = titles.index("Ссылка · 2ГИС") + 1
    for row in range(2, sheet.max_row + 1):
        cell = sheet.cell(row=row, column=column)
        if cell.value:
            assert cell.hyperlink.target.startswith("https://2gis.ru/")


def test_summary_is_honest_about_the_ceilings(report, tmp_path: Path) -> None:
    sheet = load_workbook(excel.write(report, tmp_path / "отчёт.xlsx"))["Сводка"]
    text = "\n".join(str(value) for row in sheet.iter_rows(values_only=True) for value in row if value)
    assert "кофейня" in text
    assert "Тюмень" in text
    assert "карты показывают не всё" in text
    assert "нет сайта" in text  # разбор пробелов попал в сводку


def test_default_name_says_what_and_where(report) -> None:
    assert excel.default_name(report) == "Кофейня — Тюмень — 2026-10-25.xlsx"


def test_default_name_survives_forbidden_characters(tyumen) -> None:
    result = sample_result(tyumen)
    result.settings.query = "кафе/бар: кофе?"
    assert "/" not in excel.default_name(result)
    assert ":" not in excel.default_name(result)


def test_csv_opens_in_russian_excel(report, tmp_path: Path) -> None:
    path = csvout.write(report, tmp_path / "фирмы.csv")
    raw = path.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")  # BOM: иначе кириллица превращается в кашу
    lines = path.read_text(encoding="utf-8-sig").splitlines()
    assert lines[0].split(";")[0] == "Фирма"
    assert len(lines) == len(report.places) + 1
    assert "ola@mail.ru" in lines[1] or "ola@mail.ru" in lines[2]
